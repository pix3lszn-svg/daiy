#!/usr/bin/env python3
"""Generate a clean .docx version of the management team summary."""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

ACCENT = RGBColor(0x0B, 0x6E, 0x4F)
MUTED = RGBColor(0x5A, 0x5A, 0x5A)
INK = RGBColor(0x1A, 0x1A, 0x1A)

doc = Document()
n = doc.styles["Normal"]; n.font.name = "Calibri"; n.font.size = Pt(11)
n.paragraph_format.space_after = Pt(6)
for lvl, sz in [("Heading 1", 17), ("Heading 2", 13)]:
    s = doc.styles[lvl]; s.font.name = "Calibri"; s.font.size = Pt(sz)
    s.font.bold = True; s.font.color.rgb = ACCENT if lvl == "Heading 1" else INK


def para(text, size=11, bold=False, italic=False, color=INK, after=6):
    pp = doc.add_paragraph(); pp.paragraph_format.space_after = Pt(after)
    r = pp.add_run(text); r.font.size = Pt(size); r.font.bold = bold
    r.font.italic = italic; r.font.color.rgb = color
    return pp


def b(text, lead=None):
    pp = doc.add_paragraph(style="List Bullet"); pp.paragraph_format.space_after = Pt(3)
    if lead:
        r = pp.add_run(lead); r.font.bold = True
        pp.add_run(text)
    else:
        pp.add_run(text)


def num(text, lead=None):
    pp = doc.add_paragraph(style="List Number"); pp.paragraph_format.space_after = Pt(3)
    if lead:
        r = pp.add_run(lead); r.font.bold = True
        pp.add_run(text)
    else:
        pp.add_run(text)


r = doc.add_paragraph().add_run("THE DAILY EMERALD")
r.font.size = Pt(13); r.font.bold = True; r.font.color.rgb = ACCENT
para("Heads up before the rest of the team sees this", size=22, bold=True, after=2)
para("A quick preview for the management team. Operating structure.", size=12, italic=True, color=MUTED, after=12)

para("Hey, all. Before I take this to the full staff, I want you, my management team, to see it first and tell me what you think. You are the people who will actually run this, so your feedback shapes it before anyone else gets a vote. Nothing here is final. Read it over, poke holes in it, and tell me what you would change.")
para("Here is the whole thing in plain terms.")

doc.add_heading("The problem we are fixing", level=2)
para("Right now everything runs through me. Every decision waits on me, which makes me the bottleneck and slows all of you down. The fix is structure: real teams, real leaders, and a clear way work moves, so the brand can grow without everything routing through one person.")

doc.add_heading("The structure", level=2)
b("Builders, Artists, Content Team, Voice Actors.", lead="Creative and Production: ")
b("Event Organisers, Broadcast Team, Moderation Team, Wiki Admin.", lead="Live Operations: ")
b("Members talk to their Team Lead, Team Leads talk to their Division Head, Division Heads talk to me. That single line means I coordinate two people, not forty.", lead="One clean chain of command: ")
b("Eight Team Leads, plus the two Division Heads above them.", lead="Every team gets its own lead: ")

doc.add_heading("Leadership requirements", level=2)
para("Anyone leading a group (the eight Team Leads and the Division Heads) needs to be 18 or older and have a working microphone, since leadership coordination happens in voice meetings.")

doc.add_heading("Open question I want your call on", level=2)
para("Two Division Heads each covering four teams might be too much. One option is to add a third Division Head and regroup the eight teams across three divisions instead of two. I have not decided this. I genuinely want your view, and if we go that way we work out the split together.")

doc.add_heading("The weekly rhythm", level=2)
para("A predictable cadence. Times are in UTC, convert to local.", after=4)
num("I set priorities and creative direction, brief the Division Heads.", lead="Start of week: ")
num("turn priorities into tasks and assign them through Team Leads.", lead="Division Heads ")
num("teams execute with visible status. I am not in the loop on individual tasks.", lead="Mid week: ")
num("anything needing me is batched by the Heads into one digest.", lead="Escalation is single threaded: ")
num("one or two all hands meetings a week keep everyone aligned.", lead="Alignment: ")
num("review feeds into next week's planning.", lead="End of week: ")

doc.add_heading("How work flows", level=2)
para("Each team gets one Discord forum channel. Every task is a post, assigned to a person and tagged: To Do, In Progress, Needs Review, Blocked, Done. Anyone can see what is in flight without asking. Leads keep their board healthy.")

doc.add_heading("Tools", level=2)
para("Discord only for everything internal, so nobody has to leave it to work. The handbook and weekly plan live in a read only channel, and there is a leadership only channel for the escalation funnel. No AI tools anywhere in the stack, and that includes things like Notion. The brand wiki is separate, public facing, for lore and brand content only, owned by Wiki Admin, and it contains no AI.")

doc.add_heading("How the staff meeting will run", level=2)
para("Everyone muted at the start, I present the plan, then the floor opens for questions and changes, and we converge on agreement before moving to work. I will present all of this as a proposal, not a decree, and tell the team their input gets folded in.")

doc.add_heading("A worked example to show it works", level=2)
para("A scripted Nether vs End exhibition match next week that ends in a staged in lore catastrophe, giving us an in universe reason to delay the Overworld Cup by a week. It flows cleanly through every team: Event Organisers plan it, Builders make the arena, Artists do promo, Content scripts the coverage and the disaster bulletin, Voice Actors record it, Broadcast casts it live, Moderation keeps order, and Wiki Admin records the aftermath. End to end, the structure carries it.")

doc.add_heading("The bigger picture", level=2)
para("There are much bigger plans for The Daily Emerald. All of them depend on getting organised first so we can move as a team instead of a queue.")

doc.add_heading("What I need from you", level=2)
para("Look this over and tell me, before I take it to everyone:", after=4)
b("Does the structure make sense, and are the teams and charters right?")
b("Should we add the third Division Head? If so, how would you split the eight teams?")
b("Anything missing, anything you would change, anything that would get in your way?")
para("Once you have weighed in, I will fold your feedback in, and then we take the polished version to the full staff. Thanks for being the first eyes on this.", after=6)

doc.save("/home/user/daiy/daily-emerald-ops/MANAGEMENT-TEAM-SUMMARY.docx")
print("Saved management summary docx")
