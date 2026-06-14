#!/usr/bin/env python3
"""Generate the editable Discussion Draft (.docx) for The Daily Emerald operations proposal.

Constraints honoured:
- No em dashes anywhere.
- No mention of rest, recovery, breaks, energy, or the founder's personal schedule.
- UTC for times with a note to convert to local.
- [Name TBD] placeholders where a person must be named.
- Internal stack is Discord only. No AI tools. Wiki is a separate public product.
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ACCENT = RGBColor(0x0B, 0x6E, 0x4F)
MUTED = RGBColor(0x5A, 0x5A, 0x5A)
INK = RGBColor(0x1A, 0x1A, 0x1A)

doc = Document()

# Base style
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)

for lvl, size in [("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12)]:
    st = doc.styles[lvl]
    st.font.name = "Calibri"
    st.font.size = Pt(size)
    st.font.color.rgb = ACCENT if lvl == "Heading 1" else INK
    st.font.bold = True


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), hexcolor)
    tcPr.append(sh)


def p(text="", size=11, bold=False, italic=False, color=INK, align=None, after=6, before=0):
    par = doc.add_paragraph()
    if align:
        par.alignment = align
    par.paragraph_format.space_after = Pt(after)
    par.paragraph_format.space_before = Pt(before)
    if text:
        r = par.add_run(text)
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = color
    return par


def bullet(text, level=0, bold_lead=None):
    par = doc.add_paragraph(style="List Bullet")
    if level:
        par.paragraph_format.left_indent = Inches(0.5 + 0.25 * level)
    par.paragraph_format.space_after = Pt(3)
    if bold_lead:
        r = par.add_run(bold_lead)
        r.font.bold = True
        r.font.size = Pt(11)
        r2 = par.add_run(text)
        r2.font.size = Pt(11)
    else:
        r = par.add_run(text)
        r.font.size = Pt(11)
    return par


def numbered(text, bold_lead=None):
    par = doc.add_paragraph(style="List Number")
    par.paragraph_format.space_after = Pt(3)
    if bold_lead:
        r = par.add_run(bold_lead); r.font.bold = True; r.font.size = Pt(11)
        r2 = par.add_run(text); r2.font.size = Pt(11)
    else:
        r = par.add_run(text); r.font.size = Pt(11)
    return par


def draft_tag():
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(8)
    r = par.add_run("DRAFT, open to change.")
    r.font.size = Pt(9)
    r.font.bold = True
    r.font.italic = True
    r.font.color.rgb = MUTED


def h1(text):
    doc.add_heading(text, level=1)
    draft_tag()


def h2(text):
    doc.add_heading(text, level=2)


def h3(text):
    doc.add_heading(text, level=3)


def callout(title, lines):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.style = "Table Grid"
    c = tbl.cell(0, 0)
    shade(c, "EAF3EF")
    par = c.paragraphs[0]
    r = par.add_run(title)
    r.font.bold = True
    r.font.color.rgb = ACCENT
    r.font.size = Pt(11)
    for ln in lines:
        cp = c.add_paragraph()
        cp.paragraph_format.space_after = Pt(2)
        rr = cp.add_run(ln)
        rr.font.size = Pt(10.5)
    doc.add_paragraph()


def hr():
    par = doc.add_paragraph()
    pPr = par._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "C8C8C8")
    pbdr.append(bottom)
    pPr.append(pbdr)


# ============================================================ COVER
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = title.add_run("THE DAILY EMERALD")
r.font.size = Pt(14); r.font.bold = True; r.font.color.rgb = ACCENT
p("Operations Proposal", size=30, bold=True, after=2)
p("Discussion Draft", size=20, bold=True, color=MUTED, after=14)
p("A proposal for the team to review, mark up, and shape.", size=13, italic=True, color=MUTED, after=2)
p("Nothing in this document is final.", size=13, bold=True, after=14)

# Opening note
callout("Please read first", [
    "This is a proposal, not a finished decision. It describes how The Daily Emerald could be organised, and it is meant to be edited.",
    "Your feedback is wanted. Every section here is open to change, and your suggestions will be taken into account and folded into the plan.",
    "This is a living document. We will mark it up together during the meeting.",
    "The final handbook and documentation will be generated from the approved version of this draft, so what we agree here is what ships.",
    "A light note on style: times are given in UTC, please convert to your local time. Where a specific person is needed, you will see a [Name TBD] placeholder.",
])

doc.add_page_break()

# ============================================================ CONTENTS
doc.add_heading("Contents", level=1)
contents = [
    "1. How we work, and the move from one person to a team",
    "2. Org structure and chain of command",
    "3. Team charters (all eight teams)",
    "4. Role descriptions",
    "5. The weekly operating rhythm",
    "6. How work flows and the task lifecycle",
    "7. Escalation and decisions",
    "8. Tools and where things live (Discord only)",
    "9. Volunteer expectations and code of conduct",
    "10. New staff onboarding",
    "11. Ready to use templates",
    "12. Worked example: Nether vs End",
    "13. Feedback and changes log",
]
for c in contents:
    bullet(c)
doc.add_page_break()

# ============================================================ 1. HOW WE WORK
h1("1. How we work, and the move from one person to a team")
p("The Daily Emerald is a satirical Minecraft news brand styled after the BBC. So far it has run on one person. That was workable when it was one person. It does not scale to a team of forty.")
p("Today, every decision waits on the founder. That makes one person the bottleneck and slows everyone down. The point of this document is to fix that by giving the brand a real structure: clear teams, clear leaders, and a clear way that work moves.")
h2("What changes")
bullet("You will have a team, and your team will have a lead.", bold_lead="Clear ownership. ")
bullet("You will know who to talk to, and you will not have to wait on the founder for day to day work.", bold_lead="A clear line. ")
bullet("Work will be visible, so progress is obvious without anyone chasing it.", bold_lead="Visible progress. ")
bullet("Titles and credit are real. The people doing the work get recognised for it.", bold_lead="Recognition. ")
h2("What stays the same")
p("The Daily Emerald keeps its voice, its standards, and its sense of humour. This is about how we coordinate, not about changing what we make.")
hr()

# ============================================================ 2. ORG STRUCTURE
h1("2. Org structure and chain of command")
p("The proposed structure is eight teams, grouped into two divisions, with one Division Head over each division, and the founder at the top.")
h2("The chain of command")
numbered("sets direction and priorities, and talks only to the Division Heads. Signs as The Daily Emerald.", bold_lead="Founder / Director ")
numbered("each run one division, translate priorities into tasks, and are the only people who talk to the founder.", bold_lead="Two Division Heads ")
numbered("one per team, the only person on their team who talks to the Division Head above them.", bold_lead="Eight Team Leads ")
numbered("do the work within their team.", bold_lead="Team members ")
p("The flow is Founder, then Division Heads, then eight Team Leads, then members. This single line up the chain is deliberate. It is what lets the founder coordinate two people instead of forty, and it is what stops the founder being the bottleneck.")

h2("Simple org chart")
chart = doc.add_table(rows=0, cols=1)
chart.alignment = WD_TABLE_ALIGNMENT.CENTER


def chart_row(text, fill, color=INK, bold=True, size=11):
    row = chart.add_row()
    cell = row.cells[0]
    shade(cell, fill)
    par = cell.paragraphs[0]
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = par.add_run(text)
    rr.font.bold = bold; rr.font.size = Pt(size); rr.font.color.rgb = color


chart.style = "Table Grid"
chart_row("Founder / Director", "0B6E4F", color=RGBColor(0xFF, 0xFF, 0xFF), size=12)
chart_row("Head of Creative and Production   |   Head of Live Operations", "EAF3EF", color=ACCENT, size=11)
chart_row("Creative and Production: Builders  .  Artists  .  Content Team  .  Voice Actors", "F2F4F3")
chart_row("Live Operations: Event Organisers  .  Broadcast Team  .  Moderation Team  .  Wiki Admin", "F2F4F3")
chart_row("Each team has one Team Lead, then its members", "FFFFFF", bold=False)
doc.add_paragraph()

h2("The two divisions")
h3("Creative and Production")
p("Led by the Head of Creative and Production. Oversees: Builders, Artists, Content Team, Voice Actors.")
h3("Live Operations")
p("Led by the Head of Live Operations. Oversees: Event Organisers, Broadcast Team, Moderation Team, Wiki Admin.")

h2("Leadership requirements")
p("All group leaders, meaning the eight Team Leads and the Division Heads above them, must meet two requirements:")
bullet("Must be at least 18 years of age.", bold_lead="18 or older. ")
bullet("Leadership coordination happens in voice meetings, so a working mic is required for these roles.", bold_lead="A working microphone. ")
p("These apply to leadership roles because of how coordination works. They are not a barrier to being a valued team member.")

callout("Open question for the meeting: a possible third Division Head", [
    "This is not decided. It is a question for the team to settle in the meeting.",
    "As proposed, two Division Heads each oversee four teams. That may be too much for two people.",
    "Option: add a third Division Head, then regroup the eight teams across three divisions instead of two.",
    "If the team agrees, we work out the exact split together. No grouping is fixed in advance.",
    "Decision needed: do we add a third Division Head, and if so, how should the eight teams be split across three divisions?",
])
hr()

# ============================================================ 3. TEAM CHARTERS
h1("3. Team charters")
p("A short charter for each of the eight teams. Each charter is a starting point and is open to change.")

teams = [
    ("Builders", "Creative and Production",
     "Design and construct the in game spaces, sets, and arenas that Daily Emerald content happens in.",
     ["Build event venues and sets to brief.", "Maintain and update existing builds.", "Hand finished builds to Event Organisers and Broadcast for use."],
     "An arena, set, or location is ready and matches the brief on schedule."),
    ("Artists", "Creative and Production",
     "Produce the visual identity and promotional artwork for the brand and its events.",
     ["Create promo graphics, thumbnails, and event art.", "Keep visuals consistent with the Daily Emerald look.", "Deliver assets to Content and Broadcast on time."],
     "On brand visuals are delivered in the right format and on schedule."),
    ("Content Team", "Creative and Production",
     "Write the scripts, bulletins, and coverage that carry the Daily Emerald voice.",
     ["Script match coverage, news bulletins, and in lore stories.", "Keep tone and continuity consistent.", "Hand finished scripts to Voice Actors and Broadcast."],
     "Scripts are accurate, on brand, and ready for recording and broadcast."),
    ("Voice Actors", "Creative and Production",
     "Perform and record voiced bulletins and characters from Content Team scripts.",
     ["Record bulletins and character voices to script.", "Deliver clean audio in the requested format.", "Re record as needed for quality."],
     "Recorded audio is clean, in character, and delivered on schedule."),
    ("Event Organisers", "Live Operations",
     "Plan and schedule events end to end, and coordinate the teams that make them happen.",
     ["Plan event timelines and run sheets.", "Schedule events in UTC with a local conversion note.", "Coordinate handoffs between Builders, Content, Broadcast, and others."],
     "Events are planned, scheduled, and every contributing team knows their part."),
    ("Broadcast Team", "Live Operations",
     "Cast events live to the audience and run the live production.",
     ["Run live casts and streams of events.", "Manage live production and presentation.", "Pull together builds, art, scripts, and audio into the live show."],
     "Events broadcast smoothly and on time to the audience."),
    ("Moderation Team", "Live Operations",
     "Keep the Discord server and live events orderly, safe, and on brand.",
     ["Moderate chat during events and day to day.", "Enforce the code of conduct fairly.", "Escalate serious issues through the Team Lead."],
     "The server stays orderly and welcoming, especially during live events."),
    ("Wiki Admin", "Live Operations",
     "Own and maintain the public facing Daily Emerald brand wiki, recording lore and brand content.",
     ["Set up and maintain the brand wiki (for example a MediaWiki such as Miraheze).",
      "Record events, lore, and in universe aftermath.",
      "Keep wiki content accurate and on brand. The wiki contains no AI."],
     "The brand wiki is current, accurate, and reflects the latest lore and events."),
]
for name, div, mission, resp, success in teams:
    h2(name)
    p(f"Division: {div}", italic=True, color=MUTED, after=4)
    p("Mission. " + mission)
    p("Responsibilities.", bold=True, after=2)
    for rr in resp:
        bullet(rr)
    p("Team Lead. [Name TBD]", after=2)
    p("What good looks like. " + success, after=10)
hr()

# ============================================================ 4. ROLE DESCRIPTIONS
h1("4. Role descriptions")

h2("Founder / Director")
p("Sets direction, priorities, and creative standards for the brand.")
bullet("Sets the week's priorities and creative direction.")
bullet("Briefs the Division Heads and talks only to them.")
bullet("Focuses on brand and production while the leadership team executes.")
bullet("Responds to escalations that the Division Heads batch and send up.")

h2("Division Head")
p("Runs one division and turns priorities into action. Must be 18 or older with a working microphone.")
bullet("Translates the founder's priorities into tasks and assigns them through Team Leads.")
bullet("Owns the how and the who for their division.")
bullet("Is the only person who talks to the founder, in both directions.")
bullet("Collects anything that needs the founder and sends it up as one batched digest.")

h2("Team Lead")
p("Leads one team. Must be 18 or older with a working microphone.")
bullet("Is the only person on the team who talks to the Division Head above them.")
bullet("Assigns and tracks the team's tasks, and keeps the team's forum board healthy.")
bullet("Removes blockers, and escalates upward only what genuinely needs the Head.")
bullet("Recognises and credits the team's work.")

h2("Team Member")
p("Does the work within a team.")
bullet("Picks up and delivers assigned tasks, keeping their status current.")
bullet("Flags blockers early to the Team Lead.")
bullet("Brings ideas and feedback to the team and to all hands meetings.")
hr()

# ============================================================ 5. WEEKLY RHYTHM
h1("5. The weekly operating rhythm")
p("This is the core of the whole model: a clean, predictable weekly cadence. Times are in UTC, please convert to your local time.")
numbered("The founder sets the week's priorities and creative direction and briefs the Division Heads. This is the main coordination window.", bold_lead="Start of week, planning. ")
numbered("Division Heads turn those priorities into tasks and assign them through Team Leads. They own the how and the who.", bold_lead="Translate priorities into tasks. ")
numbered("Teams work their tasks with visible status. Individual tasks do not route to the founder.", bold_lead="Mid week execution. ")
numbered("Anything that needs the founder is collected by the Division Heads and sent as one batched digest, so a small handful of people are involved rather than the whole server. Nothing reaches the founder except through the Heads.", bold_lead="Single threaded escalation. ")
numbered("One or two all hands meetings per week keep everyone aligned.", bold_lead="All hands alignment. ")
numbered("End of week review feeds straight into next week's planning.", bold_lead="End of week review. ")
p("Through the back half of the week, the founder focuses on brand and production while the leadership team executes. The cadence is designed so the brand keeps moving without routing every step through one person.")
hr()

# ============================================================ 6. HOW WORK FLOWS
h1("6. How work flows and the task lifecycle")
p("Each team has one forum channel in Discord. Every task is a post in that channel, assignable to a person, and tagged with a status.")
h2("Task statuses")
for st, desc in [
    ("To Do", "Defined and ready to be picked up."),
    ("In Progress", "Someone is actively working on it."),
    ("Needs Review", "Done by the doer, waiting on a check or sign off."),
    ("Blocked", "Stuck on something, needs help to move."),
    ("Done", "Complete and signed off."),
]:
    bullet(desc, bold_lead=f"{st}. ")
h2("The lifecycle")
numbered("A task is created as a post and tagged To Do.")
numbered("It is assigned to a member and moves to In Progress.")
numbered("When the work is done, it moves to Needs Review for a check or sign off.")
numbered("If it gets stuck at any point, it is tagged Blocked and the Team Lead helps clear it.")
numbered("Once reviewed and approved, it moves to Done.")
p("Team Leads keep their board healthy. Division Heads watch across their teams. The founder does not manage individual tasks.")
hr()

# ============================================================ 7. ESCALATION
h1("7. Escalation and decisions")
p("Escalation is single threaded on purpose. It is what keeps the founder from becoming the bottleneck again.")
bullet("Members raise blockers and questions to their Team Lead.")
bullet("Team Leads resolve what they can, and raise the rest to their Division Head.")
bullet("Division Heads resolve what they can, and batch anything that genuinely needs the founder into one digest.")
bullet("The founder responds to that digest. Nothing reaches the founder except through the Heads.")
h2("Who decides what")
bullet("Day to day execution decisions are made by Team Leads.")
bullet("Cross team and division level decisions are made by Division Heads.")
bullet("Direction, priorities, and brand level calls are made by the founder.")
hr()

# ============================================================ 8. TOOLS
h1("8. Tools and where things live")
p("The internal stack is Discord only, so volunteers never have to leave it to do their work. No AI bundled tools are used anywhere in the internal stack, and that includes tools such as Notion.")
h2("Discord is home base for everything internal")
bullet("Real time coordination and task execution happen here.")
bullet("One forum channel per team, where each task is a post tagged To Do, In Progress, Needs Review, Blocked, or Done, and tasks are assignable.")
bullet("The handbook and the weekly plan live in a read only channel, so there is no second internal tool.")
bullet("A leadership only channel holds the founder plus the Division Heads, for the escalation funnel.")
h2("The brand wiki is a separate product")
bullet("This is the public facing Daily Emerald wiki, for example a private or public MediaWiki such as Miraheze.")
bullet("It is used for lore and brand content, not for internal staff operations.")
bullet("It is owned and maintained by the Wiki Admin team, and it contains no AI.")
hr()

# ============================================================ 9. EXPECTATIONS AND CONDUCT
h1("9. Volunteer expectations and code of conduct")
p("Everyone here is an unpaid volunteer giving their time to something they care about. These expectations are about clarity and respect for that time, not command.")
h2("What we ask of each other")
bullet("Keep your task statuses current so others are not left guessing.")
bullet("Flag blockers early rather than going quiet.")
bullet("Communicate within the structure: talk to your Team Lead, and let escalation flow up the chain.")
bullet("Give realistic commitments, and say so if your availability changes.")
bullet("Times are in UTC. Convert to your local time, and be mindful that the team is global and asynchronous.")
h2("Code of conduct")
bullet("Treat everyone with respect. No harassment, discrimination, or personal attacks.")
bullet("Keep disagreements about ideas, not people.")
bullet("Respect the brand and its audience. Keep content on brand and in good taste.")
bullet("Follow Discord's terms and the rules of the server.")
bullet("If something serious comes up, raise it to a Team Lead or the Moderation Team.")
hr()

# ============================================================ 10. ONBOARDING
h1("10. New staff onboarding")
p("A simple path so a new volunteer can get productive quickly.")
numbered("Welcome and read the handbook in the read only channel.")
numbered("Get placed on a team and introduced to the Team Lead.")
numbered("Confirm role and, for leadership roles, the 18 or older and working microphone requirements.")
numbered("Tour of the team's forum channel and how task statuses work.")
numbered("Get assigned a first small task to learn the flow.")
numbered("Attend the next all hands meeting to meet the wider team.")
hr()

# ============================================================ 11. TEMPLATES
h1("11. Ready to use templates")
p("Copy and paste these. Adjust as needed.")

h2("Task brief template")
tb = doc.add_table(rows=0, cols=2); tb.style = "Table Grid"
for k, v in [
    ("Title", "Short, clear name of the task"),
    ("Team", "Which team owns it"),
    ("Assigned to", "[Name TBD]"),
    ("Status", "To Do / In Progress / Needs Review / Blocked / Done"),
    ("Due", "Date and time in UTC, convert to local"),
    ("What done looks like", "The clear outcome that means this is finished"),
    ("Depends on", "Any task or handoff this is waiting on"),
    ("Notes", "Anything else the doer needs"),
]:
    row = tb.add_row()
    row.cells[0].paragraphs[0].add_run(k).bold = True
    row.cells[1].paragraphs[0].add_run(v)
doc.add_paragraph()

h2("Weekly plan template")
wp = doc.add_table(rows=0, cols=2); wp.style = "Table Grid"
for k, v in [
    ("Week of", "Date, in UTC"),
    ("Priorities", "The two or three things that matter most this week"),
    ("Creative direction", "The brand and production focus"),
    ("By division", "Key tasks for Creative and Production, and for Live Operations"),
    ("Events this week", "What is happening, scheduled in UTC"),
    ("Open questions", "Anything to decide"),
]:
    row = wp.add_row()
    row.cells[0].paragraphs[0].add_run(k).bold = True
    row.cells[1].paragraphs[0].add_run(v)
doc.add_paragraph()

h2("Meeting agenda template")
ma = doc.add_table(rows=0, cols=2); ma.style = "Table Grid"
for k, v in [
    ("Date and time", "In UTC, with a local conversion note"),
    ("1. Open, muted", "Founder presents the week's plan and any new documentation"),
    ("2. Floor opens", "Questions and proposed changes"),
    ("3. Decisions", "Converge on agreement"),
    ("4. Close", "Owners and next steps, then move to work"),
]:
    row = ma.add_row()
    row.cells[0].paragraphs[0].add_run(k).bold = True
    row.cells[1].paragraphs[0].add_run(v)
doc.add_paragraph()
hr()

# ============================================================ 12. WORKED EXAMPLE
h1("12. Worked example: Nether vs End")
p("A concrete example that flows through the whole structure end to end.")
p("The event. A scripted Nether versus End exhibition match next week that ends in a scripted in lore catastrophe, for example a staged explosion during the match. The narrative purpose is to give us an in universe reason to delay the upcoming Overworld Cup tournament by a week. This is fictional entertainment content, the equivalent of writing a disaster into a TV plot.")
h2("How it flows, with handoffs and sign offs")
flow = [
    ("Event Organisers", "Plan and schedule the match in UTC, build the run sheet, and brief every contributing team. They sign off the overall plan and hand timelines to each team."),
    ("Builders", "Construct the arena to the brief. They hand the finished build to Event Organisers and Broadcast, who confirm it is ready for use."),
    ("Artists", "Produce promo graphics for the match. They deliver assets to the Content Team and Broadcast, who confirm they are on brand and in the right format."),
    ("Content Team", "Script both the match coverage and the in lore disaster bulletin. They hand the scripts to Voice Actors and Broadcast, and the Team Lead signs off continuity and tone."),
    ("Voice Actors", "Record the in lore bulletin from the script. They deliver clean audio to Broadcast, who confirm it is usable."),
    ("Broadcast Team", "Cast the match live, pulling together the arena, graphics, scripts, and audio. The Team Lead runs the live show and signs off that it aired as planned."),
    ("Moderation Team", "Keep the server orderly during the live event, enforcing the code of conduct and escalating anything serious through their Team Lead."),
    ("Wiki Admin", "Record the event and its in lore aftermath on the brand wiki, including the staged catastrophe and the delay of the Overworld Cup. They sign off that the wiki reflects the new canon."),
]
for who, what in flow:
    p(who + ".", bold=True, after=2)
    p(what, after=8)
p("Up the chain, the Event Organisers and other Team Leads report status to the Head of Live Operations and the Head of Creative and Production, who batch anything that needs the founder. The founder set the priority at start of week and does not touch the individual tasks. The structure carries the event from idea to aired to recorded, end to end.")
hr()

# ============================================================ 13. FEEDBACK LOG
h1("13. Feedback and changes log")
p("Use this table to capture feedback live during the meeting, so it can be edited into the draft afterward.")
log = doc.add_table(rows=1, cols=3); log.style = "Table Grid"
hdr = log.rows[0].cells
for i, t in enumerate(["Who raised it", "Requested change", "Decision"]):
    shade(hdr[i], "0B6E4F")
    rr = hdr[i].paragraphs[0].add_run(t)
    rr.font.bold = True; rr.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
for _ in range(8):
    row = log.add_row()
    for c in row.cells:
        c.paragraphs[0].add_run(" ")
doc.add_paragraph()
p("When the meeting is done, edit this draft to reflect what was agreed, then re-upload it and run Prompt 2 to generate the finalised documents.",
  italic=True, color=MUTED)

doc.save("/home/user/daiy/daily-emerald-ops/PROPOSAL-Daily-Emerald-Operations-DRAFT.docx")
print("Saved Discussion Draft")
